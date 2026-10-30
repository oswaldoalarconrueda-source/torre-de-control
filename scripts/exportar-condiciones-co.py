#!/usr/bin/env python3
"""Excel con todas las condiciones de ValerIA (agente de Joalix) en Colombia, para revisarlas con las asesoras logísticas.

Hojas: cómo decide el agente (en orden), reglas de las asesoras (las que están en el código de ValerIA), condiciones
por ciudad (auto_restricciones), la ⭐ de la matriz por ciudad, la matriz completa, transportadoras y guía, y puntos a
revisar. Lee las condiciones y la matriz con auto_exportar_condiciones (token de sincronización del .env; nunca se
imprime). Necesita openpyxl.

Uso: python3 scripts/exportar-condiciones-co.py "<salida.xlsx>"
"""
import json, os, re, sys, urllib.request
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TRANSP = {1: "Envía", 4: "Interrapidísimo", 6: "Servientrega", 10: "Coordinadora"}
TIPOS = {
    "anticipado": ("💳 Solo pago anticipado", "No sale sola. Aviso rojo en ValerIA; la asesora confirma el pago antes de despachar."),
    "anticipado_flete": ("💳 Anticipado + $20.000 de flete", "No sale sola. Aviso rojo; cobrar anticipado y $20.000 adicionales de flete."),
    "sin_cobertura": ("🚫 Sin cobertura", "No sale sola. Aviso rojo; buscar otra forma de envío o no despachar."),
    "flete_alto": ("💸 Flete muy alto", "No sale sola. Aviso rojo; la asesora decide si se despacha."),
    "solo_envia": ("🚚 Inter solo anticipado · Envía sí", "No sale sola. Aviso rojo; enviar por Envía o anticipado por Inter."),
    "inter_oficina": ("🏢 Solo a oficina Interrapidísimo", "Sale sola SOLO si la dirección es oficina de Inter; si no, aviso rojo."),
    "validar": ("📞 Validar con el cliente", "No sale sola. Aviso rojo; llamar al cliente antes de despachar."),
    "trato_especial": ("⭐ Trato especial: no se remisiona", "Nunca se pasa a remisión (ni sola ni con Ejecutar); la maneja una persona."),
}

COMO_DECIDE = [
    ("1", "Pedido con abono", "Cliente con ** al inicio del nombre", "ValerIA no la toca: ni remisión ni guía. Lo hace una persona en Effi.", "Código de ValerIA"),
    ("2", "Centro de costos", "Cliente con * al inicio del nombre (recuperado de LucidBot)", "Remisión por el centro Recuperado; los demás por Principal.", "Configuración de ValerIA"),
    ("3", "Producto, cantidad y precio", "Detalle de la orden en Effi (desde 29-oct)", "Se copian de la orden. Si no se pueden leer, se usan las guías anteriores; sin antecedentes → revisión.", "Código de ValerIA"),
    ("4", "Convenio de dropshipping", "Texto 'Dropshipping: XXXX' en la observación de la orden", "Sin convenio → revisión. El remitente sale del convenio (aprendido de las guías anteriores).", "Código de ValerIA"),
    ("5", "Transportadora (propuesta)", "Primero la regla de la asesora (hoja 'Reglas asesoras'); si no hay, la ⭐ de la matriz", "Es lo que sale por defecto en 'Ejecutar'; la persona puede cambiarla.", "Código de ValerIA + matriz"),
    ("6", "⭐ de la matriz", "Mayor % de entrega a esa ciudad (sin Domina ni TCC); empate → menor flete", "Ciudades sin flete en la matriz no se recomiendan. Ver hoja 'Matriz ⭐ por ciudad'.", "Matriz de entrega"),
    ("7", "Condición por ciudad", "Ciudad, departamento o palabra en la dirección (hoja 'Condiciones por ciudad')", "Si la orden tiene condición, NO sale sola; queda con aviso rojo para la asesora.", "Tabla de condiciones"),
    ("8", "Va sola en Manual (oficina)", "Dirección con OFICINA / OF / RECLAM + INTERRAPIDISIMO o SERVIENTREGA", "Remisión y guía solas por esa transportadora, sin pulsar Ejecutar.", "Código de ValerIA"),
    ("9", "Va sola en Manual (coincidencia)", "La regla de la asesora y la ⭐ de la matriz dicen la misma transportadora", "Remisión y guía solas (o solo la guía si la remisión la hizo una persona).", "Código de ValerIA"),
    ("10", "Todo lo demás en Manual", "—", "Queda en 'Órdenes sin remisión' hasta que alguien pulse 'Ejecutar' (escoge paso y transportadora).", "—"),
    ("11", "Modo Automático (si se activa)", "Toda orden lista, de a 8 por corrida", "Se crea sola, excepto las que tienen condición por ciudad o abono.", "Mando de ValerIA"),
    ("12", "Ya resuelta", "La orden ya tiene guía vigente del mismo cliente, o alguien pulsó 'Ya está resuelta'", "Sale de la lista y ValerIA no la crea.", "Base de datos"),
    ("13", "Horario", "Cada 10 minutos y 2 minutos antes de cada corte de despacho", "Corridas de ValerIA (prepara planes y crea lo que va solo).", "Horario de ValerIA"),
]

REGLAS_ASESORAS = [
    ("1", "Reclama en oficina Interrapidísimo", "Dirección con OFICINA/OF/RECLAM e INTERRAPIDISIMO", "Cualquiera", "Interrapidísimo"),
    ("2", "Reclama en oficina Servientrega", "Dirección con OFICINA/OF/RECLAM y SERVIENTREGA", "Cualquiera", "Servientrega"),
    ("3", "Medellín: barrio popular (Coordinadora genera devoluciones)", "Dirección con ENCISO, SANTO DOMINGO, BELENCITO, POPULAR, SAN JAVIER, ROBLEDO o MORAVIA",
     "Medellín, Envigado, Itagüí, Sabaneta, La Estrella, Caldas, Copacabana, Girardota, Barbosa (Antioquia)", "Envía"),
    ("4", "Medellín y Valle de Aburrá", "—", "Medellín, Envigado, Itagüí, Sabaneta, La Estrella, Caldas, Copacabana, Girardota, Barbosa (Antioquia)", "Coordinadora"),
    ("5", "Bello: Cabañas, Santa Ana, Niquía o centro", "Dirección con CABAÑAS, SANTA ANA, NIQUIA o CENTRO", "Bello", "Coordinadora"),
    ("6", "Bello (resto)", "—", "Bello", "Envía"),
    ("7", "Cali", "—", "Cali", "Coordinadora"),
    ("8", "Ciudades por Envía", "—", "Barranquilla, Santa Marta, Cúcuta, Montería, Bogotá", "Envía"),
    ("9", "Ciudades por Servientrega", "—", "Florencia, Barrancabermeja, Manizales, Pasto, Popayán, Neiva, Ibagué, Bucaramanga", "Servientrega"),
    ("10", "Tolima (resto)", "—", "Departamento del Tolima", "Envía"),
    ("11", "Pueblo o zona alejada (solo llega Inter)", "La matriz solo tiene datos de Inter para esa ciudad", "Cualquiera", "Interrapidísimo (no cuenta como coincidencia)"),
    ("12", "Atlántico, Magdalena o Buenaventura", "—", "Atlántico, Magdalena, Buenaventura", "Interrapidísimo"),
    ("—", "Sin regla", "—", "El resto del país", "Se propone la ⭐ de la matriz"),
]

GUIA = [
    ("Transportadoras habilitadas", "Envía (1), Interrapidísimo (4), Servientrega (6), Coordinadora (10). Domina y TCC están excluidas."),
    ("Valor declarado · Envía", "$10.000 fijo"),
    ("Valor declarado · Interrapidísimo", "El valor del recaudo, mínimo $45.000"),
    ("Valor declarado · Servientrega", "El valor del recaudo"),
    ("Valor declarado · Coordinadora", "El valor del pedido, máximo $100.000"),
    ("Paquete", "1 paquete, 1 kg, volumen 1"),
    ("Medidas", "Servientrega y Coordinadora: 1 × 1 × 1 cm (Effi las exige)"),
    ("Contenido de la guía", "Se copia de la remisión"),
    ("Origen", "La mercancía sale siempre de Bogotá"),
]

REVISAR = [
    ("La matriz no trae departamento", "Pueblos con el mismo nombre en varios departamentos (p. ej. San Pablo, Bolívar) comparten fila; se dejó el dato más prudente (menor % de entrega). Puede proponer una transportadora rara en esos pueblos."),
    ("Detección de oficina", "Solo se reconoce si la dirección dice OFICINA/OF/RECLAM y además INTERRAPIDISIMO o SERVIENTREGA. 'Reclama en Inter' o 'oficina Inter' (sin 'rapidísimo') no van solas."),
    ("Nombres de ciudad", "Las condiciones por ciudad se cruzan por el nombre exacto (sin tildes ni espacios). Una ciudad escrita distinto en Effi no toma la condición: revisar la columna 'Claves'."),
    ("Regla de asesora vs. matriz", "La regla de la asesora manda sobre la ⭐. Si la matriz muestra mejor entrega por otra transportadora (p. ej. Bogotá), la propuesta sigue siendo la de la asesora."),
    ("Palabras de barrio (Medellín y Bello)", "Se buscan como palabras en la dirección; 'Centro' en Bello o 'Popular' en Medellín pueden aparecer en direcciones que no son de ese barrio."),
]


def env():
    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    vals = {}
    for linea in open(os.path.join(raiz, ".env")):
        m = re.match(r"^([A-Z_]+)=(.*)$", linea.strip())
        if m:
            vals[m.group(1)] = m.group(2).strip().strip('"')
    return vals["NEXT_PUBLIC_SUPABASE_URL"], vals["NEXT_PUBLIC_SUPABASE_ANON_KEY"], vals["EFFI_SYNC_TOKEN"]


def traer():
    url, llave, token = env()
    req = urllib.request.Request(url.rstrip("/") + "/rest/v1/rpc/auto_exportar_condiciones", method="POST",
        data=json.dumps({"p_token": token, "p_cuenta": "TIENDA_CO", "p_pais": "Colombia"}).encode(),
        headers={"apikey": llave, "Authorization": "Bearer " + llave, "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read())


AZUL = PatternFill("solid", fgColor="1F3A68")


def hoja(wb, titulo, encabezado, filas, anchos):
    ws = wb.create_sheet(titulo)
    ws.append(encabezado)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = AZUL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(list(f))
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main(salida):
    d = traer()
    wb = Workbook()
    wb.remove(wb.active)
    hoja(wb, "Cómo decide", ["Paso", "Regla", "Cuándo aplica", "Qué hace ValerIA", "Dónde está"], COMO_DECIDE, [6, 28, 48, 62, 24])
    hoja(wb, "Reglas asesoras", ["Prioridad", "Regla", "Palabras en la dirección", "Ciudad / zona", "Transportadora"], REGLAS_ASESORAS, [10, 44, 52, 52, 30])

    cond = []
    for r in d["restricciones"]:
        etiqueta, efecto = TIPOS.get(r["tipo"], (r["tipo"], ""))
        cond.append((r.get("dpto") or "Cualquiera", r["lugar"], ", ".join(r.get("ciudades") or []), r.get("palabra") or "",
                     ", ".join(r.get("en_ciudades") or []), etiqueta, r["motivo"], efecto))
    hoja(wb, "Condiciones por ciudad", ["Departamento (clave)", "Lugar", "Claves de ciudad", "Palabra en la dirección", "Solo dentro de",
                                         "Condición", "Motivo de la asesora", "Efecto en ValerIA"], cond, [20, 34, 34, 22, 18, 30, 40, 52])

    por_ciudad = defaultdict(list)
    for m in d["matriz"]:
        por_ciudad[m["destino"]].append(m)
    estrella = []
    for ciudad, filas in sorted(por_ciudad.items(), key=lambda x: x[0]):
        con = sorted([f for f in filas if f["flete_entregado"] is not None], key=lambda f: (-(f["entrega"] or 0), f["flete_entregado"]))
        mejor = con[0] if con else None
        solo_inter = bool(con) and all(f["transportadora"] == 4 for f in con)
        fila = [ciudad, TRANSP.get(mejor["transportadora"], mejor["nombre"]) if mejor else "Sin datos",
                mejor["entrega"] if mejor else None, mejor["flete_entregado"] if mejor else None, "Sí" if solo_inter else ""]
        for t in (1, 4, 6, 10):
            f = next((x for x in filas if x["transportadora"] == t), None)
            fila += [f["entrega"] if f and f["flete_entregado"] is not None else None, f["flete_entregado"] if f else None]
        estrella.append(fila)
    enc = ["Ciudad", "⭐ Recomendada", "% entrega ⭐", "Flete ⭐", "Solo llega Inter"]
    for t in (1, 4, 6, 10):
        enc += [f"% entrega {TRANSP[t]}", f"Flete {TRANSP[t]}"]
    ws = hoja(wb, "Matriz ⭐ por ciudad", enc, estrella, [26, 18, 12, 12, 10] + [12, 12] * 4)
    for fila in ws.iter_rows(min_row=2):
        for c in fila[2:]:
            if isinstance(c.value, (int, float)):
                c.number_format = '0.0' if "entrega" in str(ws.cell(1, c.column).value) else '"$"#,##0'

    completa = [(m["destino"], TRANSP.get(m["transportadora"], m["nombre"]), m["entrega"], m["devol"], m["exposicion"],
                 m["flete_entregado"], m["flete_devolucion"]) for m in d["matriz"]]
    hoja(wb, "Matriz completa", ["Ciudad", "Transportadora", "% entrega", "% devolución", "% exposición", "Flete entregado", "Flete devolución"],
         completa, [26, 18, 12, 12, 12, 14, 14])
    hoja(wb, "Transportadoras y guía", ["Tema", "Regla"], GUIA, [34, 90])
    hoja(wb, "Para revisar", ["Punto", "Por qué puede verse raro"], REVISAR, [34, 110])
    wb.save(salida)
    print(json.dumps({"condiciones": len(cond), "ciudades_matriz": len(estrella), "filas_matriz": len(completa)}, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])
